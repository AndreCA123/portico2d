"""Pasos del método en HTML (para la ventana y el PDF) y exportación a Excel."""
from __future__ import annotations
import numpy as np

from .solver import Resultado, esfuerzos

AZUL = "#1d2b53"
ROJO = "#c1121f"


def _num(v, dec):
    if abs(v) < 0.5 * 10 ** (-dec):
        return "0"
    return f"{v:.{dec}f}".replace("-", "−")


import re
import html as _html

_MODO = {"pre": False, "pt": 9.5}   # True → matrices en texto monoespaciado (PDF)


def _plano(t):
    return _html.unescape(re.sub(r"<[^>]+>", "", t))


class Bloque:
    """Bloque de texto monoespaciado: lista de (texto_plano, html)."""

    def __init__(self, lineas):
        self.lineas = lineas
        self.ancho = max((len(p) for p, _ in lineas), default=0)

    @staticmethod
    def texto(t):
        return Bloque([(_plano(t), t)])

    def __radd__(self, otro):
        return combinar([Bloque.texto(otro) if isinstance(otro, str) else otro, self])

    def __add__(self, otro):
        return combinar([self, Bloque.texto(otro) if isinstance(otro, str) else otro])

    def __str__(self):
        cuerpo = "\n".join(h for _, h in self.lineas)
        pt = _MODO["pt"]
        return ('<pre style="font-family:DejaVu Sans Mono,monospace;font-size:%spt">' % pt) + cuerpo + "</pre>"


def combinar(bloques, sep=" "):
    alto = max(len(b.lineas) for b in bloques)
    cols = []
    for b in bloques:
        pad = (alto - len(b.lineas)) // 2
        filas = [("", "")] * pad + b.lineas + [("", "")] * (alto - len(b.lineas) - pad)
        cols.append([(p + " " * (b.ancho - len(p)), h + " " * (b.ancho - len(p))) for p, h in filas])
    lineas = []
    for k in range(alto):
        lineas.append((sep.join(c[k][0] for c in cols), sep.join(c[k][1] for c in cols)))
    return Bloque(lineas)


def _mat_pre(Mx, dec, filas, cols):
    nr, nc = Mx.shape
    txt = [[_num(v, dec) for v in fila_] for fila_ in Mx]
    w = [max(len(txt[i][j]) for i in range(nr)) for j in range(nc)]
    if cols:
        w = [max(w[j], len(cols[j])) for j in range(nc)]
    wf = max((len(f) for f in filas), default=0) if filas else 0
    R = lambda c: f'<font color="{ROJO}">{c}</font>'
    lineas = []
    if cols:
        cab = " " * (wf + 2) + " ".join(c.rjust(w[j]) for j, c in enumerate(cols))
        lineas.append((cab + "  ", f'<font color="#888888">{_html.escape(cab)}</font>  '))
    for i in range(nr):
        if nr == 1:
            a, b = "[", "]"
        elif i == 0:
            a, b = "⎡", "⎤"
        elif i == nr - 1:
            a, b = "⎣", "⎦"
        else:
            a, b = "⎢", "⎥"
        f = (filas[i] if filas else "").rjust(wf)
        cuerpo = " ".join(txt[i][j].rjust(w[j]) for j in range(nc))
        plano = f"{f} {a}{cuerpo} {b}"
        lineas.append((plano, f'<font color="#888888">{_html.escape(f)}</font> {R(a)}{cuerpo} {R(b)}'))
    return Bloque(lineas)


def _mat_katex(Mx, dec, filas, cols):
    """Matriz en LaTeX (para KaTeX en la versión web): corchetes reales, filas y columnas rotuladas."""
    nr, nc = Mx.shape
    txt = [[_num(v, dec).replace("−", "-") for v in fila_] for fila_ in Mx]
    cuerpo = r" \\ ".join(" & ".join(f) for f in txt)
    m = r"\left[\begin{array}{" + "r" * nc + "}" + cuerpo + r"\end{array}\right]"
    if filas or cols:
        lab = r"\begin{array}{r}" + r" \\ ".join(
            r"\scriptstyle\mathrm{" + (filas[i] if filas else "") + "}" for i in range(nr)) + r"\end{array}"
        if cols:
            celdas = []
            for j in range(nc):
                ancho = max((txt[i][j] for i in range(nr)), key=len)
                extra = r"\hphantom{" + cols[j] + "}" if len(cols[j]) > len(ancho) else ""
                celdas.append(r"\hphantom{" + ancho + "}" + extra + r"\mathllap{\scriptstyle\mathrm{" + cols[j] + "}}")
            cab = (r"\hphantom{[}" + (r"\kern0.14em " if nr > 1 else "") +
                   r"\begin{array}{" + "r" * nc + "}" + " & ".join(celdas) + r"\end{array}")
        else:
            cab = ""
        m = r"\begin{array}{rl} & " + cab + r" \\ " + lab + " & " + m + r" \end{array}"
    return '<span class="katex-math">' + _html.escape(m, quote=False) + "</span>"


def mat(Mx, dec=2, filas=None, cols=None):
    """Matriz con corchetes (barras rojas) compatible con QTextDocument."""
    Mx = np.atleast_2d(np.asarray(Mx, dtype=float))
    if Mx.shape[0] == 1 and (filas is None or len(filas) > 1):
        Mx = Mx.T
    if _MODO.get("katex"):
        return _mat_katex(Mx, dec, filas, cols)
    if _MODO["pre"]:
        return _mat_pre(Mx, dec, filas, cols)
    nr, nc = Mx.shape
    h = '<table cellspacing="0" cellpadding="2" style="margin:2px">'
    if cols:
        h += "<tr><td></td><td></td>" + "".join(
            f'<td align="center"><font color="#888" size="2">{c}</font></td>' for c in cols) + "<td></td></tr>"
    for i in range(nr):
        h += "<tr>"
        h += f'<td align="right"><font color="#888" size="2">{filas[i] if filas else ""}</font></td>'
        h += f'<td width="3" bgcolor="{ROJO}"></td>'
        h += "".join(f'<td align="right" nowrap style="padding-left:6px;padding-right:6px">{_num(v, dec)}</td>'
                     for v in Mx[i])
        h += f'<td width="3" bgcolor="{ROJO}"></td></tr>'
    return h + "</table>"


def fila(*partes):
    if _MODO.get("katex"):
        return '<div class="mrow">' + "".join(f'<div class="mi">{p}</div>' for p in partes) + "</div>"
    if _MODO["pre"]:
        return str(combinar([p if isinstance(p, Bloque) else Bloque.texto(p) for p in partes], sep="  "))
    celdas = "".join(f'<td valign="middle" nowrap>{p}</td>' for p in partes)
    return f'<table cellspacing="4"><tr>{celdas}</tr></table>'


def paso(n, titulo):
    return f'<h3 style="color:{AZUL}"><font color="{ROJO}">Paso {n}</font> · {titulo}</h3>'


def html_pasos(res: Resultado, imagen_modelo: str | None = None, pdf: bool = False,
               figuras: dict | None = None, katex: bool = False) -> str:
    """figuras: {'numeracion': url, 'configuraciones': [url, ...], 'fijacion': url} (ver figuras.py).
    katex=True (versión web): las matrices salen como <span class="katex-math">LaTeX</span>."""
    _MODO["pre"] = True
    _MODO["katex"] = bool(katex)
    _MODO["pt"] = 7.6 if pdf else 10
    _MODO["figs"] = figuras or {}
    try:
        return _html_pasos(res, imagen_modelo, pdf)
    finally:
        _MODO["pre"] = False
        _MODO["katex"] = False


def _html_pasos(res, imagen_modelo, pdf):
    m = res.modelo
    nD = len(res.D)
    Dn = [f"D{j + 1}" for j in range(nD)]
    h = [f'<h2 style="color:{AZUL}">{m.titulo}</h2>',
         f"<p>Unidades: t, m, rad. EA = {'∞ (sin deformación axial)' if m.ea_infinito else 'finito'}.</p>"]
    if imagen_modelo:
        h.append(f'<p align="center"><img src="{imagen_modelo}" width="560"></p>')
    h.append(paso(1, "Sistema de referencia global"))
    h.append("<p>X → derecha (+), Y ↑ arriba (+), giro antihorario (+).</p>")
    h.append(paso(2, "Nudos y elementos"))
    t = '<table border="1" cellspacing="0" cellpadding="3"><tr><th>Nudo</th><th>X</th><th>Y</th><th>Apoyo</th></tr>'
    for n in m.nudos:
        t += f"<tr><td>{n.id}</td><td>{n.x:g}</td><td>{n.y:g}</td><td>{n.apoyo}</td></tr>"
    t += "</table>"
    hay_rot = any(rb.rot_i or rb.rot_j for rb in res.barras)
    hay_br = any(rb.bi or rb.bj for rb in res.barras)
    t2 = ('<table border="1" cellspacing="0" cellpadding="3"><tr><th>Barra</th><th>i → j</th>'
          '<th>L</th><th>cos α</th><th>sen α</th><th>EI</th>' + ("" if m.ea_infinito else "<th>EA</th>")
          + ("<th>Extremos</th>" if hay_rot else "")
          + ("<th>brazo i</th><th>brazo j</th><th>L flexible</th>" if hay_br else "") + "</tr>")
    for rb in res.barras:
        g = rb.geo
        ext = {(False, False): "rígido–rígido", (True, False): "rótula en i", (False, True): "rótula en j",
               (True, True): "rótula i y j"}[(bool(rb.rot_i), bool(rb.rot_j))]
        t2 += (f"<tr><td>{rb.id}</td><td>{rb.i} → {rb.j}</td><td>{g.L:.4g}</td><td>{g.c:.4f}</td>"
               f"<td>{g.s:.4f}</td><td>{rb.EI:g}</td>" + ("" if m.ea_infinito else f"<td>{rb.EA:g}</td>")
               + (f"<td>{ext}</td>" if hay_rot else "")
               + (f"<td>{rb.bi:g}</td><td>{rb.bj:g}</td><td>{rb.Lf:.4g}</td>" if hay_br else "") + "</tr>")
    t2 += "</table>"
    if _MODO.get("katex"):
        h.append('<div class="table-responsive">' + t + '</div>')
        h.append('<div class="table-responsive" style="margin-top: 10px;">' + t2 + '</div>')
    else:
        h.append(fila(t, "&nbsp;&nbsp;", t2))
    h.append(paso(3, "Sistemas de referencia locales"))
    h.append("<p>Reglas: (1) el eje x local es paralelo al elemento; (2) su sentido + va del nudo de "
             "<b>menor</b> número al de <b>mayor</b> número (i &lt; j); (3) el eje z local tiene el mismo "
             "sentido que el z global (sale del plano); (4) el eje y local sigue la regla de la mano derecha "
             "(y = z × x: x girado 90° antihorario).</p>")
    h.append("<p>" + " ; ".join(f"Barra {rb.id}: {rb.i} → {rb.j}" for rb in res.barras) + "</p>")
    h.append(paso(4, "GDL globales"))
    h.append("<p>Orden: primero los desplazamientos paralelos a X, luego los paralelos a Y y al final "
             "los giros; dentro de cada grupo, por número de nudo.</p>")
    h.append("<ul>" + "".join(f"<li><b>D{j + 1}</b> = {e}</li>" for j, e in enumerate(res.gdl.etiquetas)) + "</ul>")
    if m.ea_infinito and res.gdl.dependientes:
        h.append("<p>Por EA = ∞ (las barras no cambian de longitud): " +
                 "; ".join(res.gdl.dependientes) + "</p>")
    h.append(paso(5, "GDL locales"))
    h.append("<p>{d} = [" + ", ".join(res.barras[0].filas) + "]" +
             (" (sin axial por EA = ∞)" if m.ea_infinito else "") + "</p>")
    figs = _MODO.get("figs") or {}
    ancho = 520 if pdf else 640
    if figs.get("numeracion"):
        h.append(f'<p align="center"><img src="{figs["numeracion"]}" width="{ancho}"></p>')
    h.append(paso(6, "Matrices [A]   ({d} = [A]{D})"))
    if figs.get("configuraciones"):
        h.append("<p>Configuraciones deformadas vinculadas a cada GDL global (D<sub>j</sub> = 1 y los demás = 0):</p>")
        cfg = figs["configuraciones"]
        w = 300 if pdf else 380
        t = '<table align="center" cellspacing="4">'
        for k in range(0, len(cfg), 2):
            t += "<tr>" + "".join(f'<td><img src="{u}" width="{w}"></td>' for u in cfg[k:k + 2]) + "</tr>"
        h.append(t + "</table>")
    h.append("<p>Columna j = desplazamientos locales cuando D<sub>j</sub> = 1. "
             "Transversal local: v = −sen α · u + cos α · v.</p>")
    if hay_br:
        h.append("<p>Con brazos rígidos, {d} son los desplazamientos en las caras de la parte flexible: "
                 "v'<sub>i</sub> = v<sub>i</sub> + b<sub>i</sub>·θ<sub>i</sub> ; "
                 "v'<sub>j</sub> = v<sub>j</sub> − b<sub>j</sub>·θ<sub>j</sub> (θ no cambia).</p>")
    por_fila = 2 if pdf else 3
    partes = [f"<b>[A]{rb.id}</b>" + mat(rb.A, 4, rb.filas, Dn) for rb in res.barras]
    for k in range(0, len(partes), por_fila):
        h.append(fila(*partes[k:k + por_fila]))
    h.append(paso(7, "Matrices [k] locales"))
    if hay_rot:
        h.append("<p>Con rótula, [k] se condensa: [k]c = k<sub>aa</sub> − k<sub>ab</sub>·k<sub>bb</sub>⁻¹·k<sub>ba</sub> "
                 "(fila y columna del giro liberado = 0; con una rótula aparecen 3EI/L, 3EI/L², 3EI/L³). "
                 "Con rótula en ambos extremos solo queda la rigidez axial (armadura).</p>")
    for rb in res.barras:
        L = rb.Lf
        h.append(f"<p><b>Barra {rb.id}</b> ({'L flexible' if (rb.bi or rb.bj) else 'L'} = {L:.4g}): 12EI/L³ = {12 * rb.EI / L**3:.4g} ; "
                 f"6EI/L² = {6 * rb.EI / L**2:.4g} ; 4EI/L = {4 * rb.EI / L:.4g} ; 2EI/L = {2 * rb.EI / L:.4g}"
                 + ("" if m.ea_infinito else f" ; EA/L = {rb.EA / L:.4g}") + "</p>")
        h.append(str(mat(rb.k, 2, rb.filas, rb.filas)))
    h.append(paso(8, "[K]e = [A]ᵀ [k] [A]"))
    partes = [f"<b>[K]{rb.id}</b>" + mat(rb.Kg, 2, Dn, Dn) for rb in res.barras]
    for k in range(0, len(partes), por_fila):
        h.append(fila(*partes[k:k + por_fila]))
    h.append(paso(9, "Ensamble [K] = Σ [K]e"))
    sim = np.allclose(res.K, res.K.T)
    h.append(fila("<b>[K] =</b>", mat(res.K, 2, Dn, Dn),
                  f"&nbsp; simétrica: {'sí ✓' if sim else 'NO'} ; det = {np.linalg.det(res.K):.4e}"))
    h.append(paso(10, "Cargas de fijación {r} (ejes locales)"))
    h.append("<p>{r} = fuerzas del empotramiento sobre la barra (V en +y local, M antihorario +).</p>")
    if figs.get("fijacion"):
        h.append(f'<p align="center"><img src="{figs["fijacion"]}" width="{ancho}"></p>')
    if hay_rot:
        h.append("<p>Barras con rótula: {r}c = r<sub>a</sub> − k<sub>ab</sub>·k<sub>bb</sub>⁻¹·r<sub>b</sub> "
                 "(momento nulo en la rótula).</p>")
    partes = []
    for rb in res.barras:
        if rb.usa_axial_en_R:
            partes.append(f"<b>{{r}}{rb.id}</b>" + mat(rb.r6, 3, ["N_i", "V_i", "M_i", "N_j", "V_j", "M_j"]))
        else:
            partes.append(f"<b>{{r}}{rb.id}</b>" + mat(rb.r, 3, rb.filas))
    for k in range(0, len(partes), 4 if pdf else 6):
        h.append(fila(*partes[k:k + (4 if pdf else 6)]))
    if any(rb.usa_axial_en_R for rb in res.barras):
        h.append("<p><i>Las barras con carga axial usan [A] completa (6 filas) para pasar esa carga a los D.</i></p>")
    for rb in res.barras:
        if np.any(np.abs(rb.r_brazos) > 1e-12):
            h.append(f"<p><i>Barra {rb.id}: cargas sobre los brazos rígidos (pasan directo a los nudos; "
                     "se suman a {R} como [A<sub>nudo</sub>]ᵀ{r}<sub>brazo</sub>):</i></p>")
            h.append(fila("{r}<sub>brazo</sub> =", mat(rb.r_brazos, 3, ["N_i", "V_i", "M_i", "N_j", "V_j", "M_j"])))
    h.append(paso(11, "{R} = Σ [A]ᵀ{r} − {fuerzas aplicadas directamente en los nudos}"))
    partes = []
    for k, rb in enumerate(res.barras):
        partes += (["+"] if k else []) + [mat(rb.Re, 3, Dn if k == 0 else None)]
    if pdf and len(res.barras) > 2:
        h.append(fila("<b>Σ [A]ᵀ{r} :</b>", *partes))
        h.append(fila("<b>Σ [A]ᵀ{r} =</b>", mat(res.Rs, 3, Dn)))
    else:
        h.append(fila("<b>Σ [A]ᵀ{r} =</b>", *partes, "=", mat(res.Rs, 3)))
    h.append("<p>Fuerzas aplicadas directamente en los nudos, en la dirección de cada D:</p>")
    h.append(fila("<b>{R} =</b>", mat(res.Rs, 3, Dn), "−", mat(res.Q, 3), "=", mat(res.R, 3)))
    h.append(paso(12, "{F} = −{R}"))
    h.append("<p>Como [K]{D} + {R} = 0 → {F} = −{R} y [K]{D} = {F}.</p>")
    h.append(fila("<b>{F}</b> =", mat(res.F, 3, Dn)))
    h.append(paso(13, "Desplazamientos {D} = [K]⁻¹{F}"))
    if pdf:
        h.append(fila("<b>[K]⁻¹ =</b>", mat(np.linalg.inv(res.K), 7)))
        h.append(fila("<b>{D} =</b>", mat(res.D, 6, Dn)))
    else:
        h.append(fila("<b>[K]⁻¹ =</b>", mat(np.linalg.inv(res.K), 7), "<b>{D} =</b>", mat(res.D, 6, Dn)))
    h.append(paso(14, "Fuerzas internas: {d} = [A]{D} ; {q}c = [k]{d} ; {q} = {q}c + {r}"))
    for rb in res.barras:
        h.append(f"<p><b>Barra {rb.id}</b></p>")
        if pdf:
            h.append(fila("{d} =", mat(rb.d, 6, rb.filas), "{q}c = [k]{d} =", mat(rb.qc, 3)))
            h.append(fila("{q} = {q}c + {r} =", mat(rb.qc, 3), "+", mat(rb.r, 3), "=", mat(rb.q, 2)))
        else:
            h.append(fila("{d} =", mat(rb.d, 6, rb.filas), "{q}c =", mat(rb.qc, 3), "+", mat(rb.r, 3),
                          "= <b>{q}</b> =", mat(rb.q, 2)))
    # resumen
    h.append(f'<h3 style="color:{AZUL}">Resumen</h3>')
    t = ('<table border="1" cellspacing="0" cellpadding="3"><tr><th>Barra</th><th>N (t)</th>'
         '<th>V_i</th><th>M_i</th><th>V_j</th><th>M_j</th><th>M máx (+)</th><th>M mín (−)</th></tr>')
    for rb in res.barras:
        xs, N, V, M, _, _ = esfuerzos(rb)
        t += (f"<tr><td>{rb.id}</td><td>{_num(N[0], 2)}</td><td>{_num(rb.q6[1], 2)}</td><td>{_num(rb.q6[2], 2)}</td>"
              f"<td>{_num(rb.q6[4], 2)}</td><td>{_num(rb.q6[5], 2)}</td>"
              f"<td>{_num(M.max(), 2)} (x={xs[M.argmax()]:.2f})</td><td>{_num(M.min(), 2)} (x={xs[M.argmin()]:.2f})</td></tr>")
    t += "</table>"
    h.append(t)
    h.append("<p><i>q = fuerzas en extremos (ejes locales). En los diagramas M(x) = −M_i + V_i·x − cargas; "
             "M &gt; 0 tracciona la cara −y local. N &lt; 0 = compresión"
             + (" (con EA = ∞, N sale del equilibrio)." if m.ea_infinito else ".")
             + (" Con brazos rígidos, V y M de la tabla son los de las caras de la parte flexible." if hay_br else "")
             + "</i></p>")
    t = '<table border="1" cellspacing="0" cellpadding="3"><tr><th>Apoyo</th><th>Rx (t)</th><th>Ry (t)</th><th>M (t·m)</th></tr>'
    for nid, (rx, ry, mz) in res.reacciones.items():
        t += f"<tr><td>Nudo {nid}</td><td>{_num(rx, 2)}</td><td>{_num(ry, 2)}</td><td>{_num(mz, 2)}</td></tr>"
    h.append("<p><b>Reacciones</b></p>" + t + "</table>")
    t = '<table border="1" cellspacing="0" cellpadding="3"><tr><th>Nudo</th><th>u (m)</th><th>v (m)</th><th>θ (rad)</th></tr>'
    for nid, (ux, uy, th) in res.u_nudos.items():
        t += f"<tr><td>{nid}</td><td>{_num(ux, 6)}</td><td>{_num(uy, 6)}</td><td>{_num(th, 6)}</td></tr>"
    h.append("<p><b>Desplazamientos de nudos</b></p>" + t + "</table>")
    for a in res.avisos:
        h.append(f'<p><font color="{ROJO}">{a}</font></p>')
    return "\n".join(h)


# ----------------------------------------------------------------- Excel
def exportar_excel(res: Resultado, ruta: str, imagenes: dict | None = None, figuras: dict | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.drawing.image import Image as XImg

    wb = Workbook()
    ws = wb.active
    ws.title = "Pasos"
    B = Font(bold=True)
    T = Font(bold=True, size=13, color="1D2B53")
    F = PatternFill("solid", fgColor="DDEBF7")
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    r = [1]

    def linea(txt, font=None):
        c = ws.cell(r[0], 1, txt)
        if font:
            c.font = font
        r[0] += 1

    def matriz(nombre, Mx, fmt="0.0000", filas=None, cols=None):
        Mx = np.atleast_2d(np.asarray(Mx, float))
        if Mx.shape[0] == 1:
            Mx = Mx.T
        ws.cell(r[0], 1, nombre).font = B
        if cols:
            for j, cn in enumerate(cols):
                ws.cell(r[0], 3 + j, cn).font = Font(italic=True, color="888888")
            r[0] += 1
        for i in range(Mx.shape[0]):
            if filas:
                ws.cell(r[0], 2, filas[i]).font = Font(italic=True, color="888888")
            for j in range(Mx.shape[1]):
                c = ws.cell(r[0], 3 + j, float(Mx[i, j]))
                c.number_format = fmt
                c.fill = F
                c.border = box
            r[0] += 1
        r[0] += 1

    m = res.modelo
    nD = len(res.D)
    Dn = [f"D{j + 1}" for j in range(nD)]
    linea(m.titulo, T)
    linea("Unidades: t, m, rad. Convención: X→+, Y↑+, antihorario +.")
    linea("SR local: x paralelo a la barra, de nudo menor a mayor (i < j); z local = z global; "
          "y por la mano derecha (x girado 90° antihorario).")
    linea("Barras: " + " ; ".join(f"{rb.id}: {rb.i}→{rb.j}" for rb in res.barras))
    linea("GDL globales (orden: desplazamientos en X, luego en Y, luego giros; por número de nudo):", B)
    for j, e in enumerate(res.gdl.etiquetas):
        linea(f"  D{j + 1} = {e}")
    if res.gdl.dependientes:
        linea("  " + "; ".join(res.gdl.dependientes))
    for rb in res.barras:
        extra = []
        if rb.rot_i or rb.rot_j:
            extra.append("rótula en " + " y ".join(e for e, f in (("i", rb.rot_i), ("j", rb.rot_j)) if f)
                         + " ([k] y {r} condensadas)")
        if rb.bi or rb.bj:
            extra.append(f"brazos rígidos bi = {rb.bi:g} m, bj = {rb.bj:g} m, L flexible = {rb.Lf:.4g} m")
        if extra:
            linea(f"  Barra {rb.id}: " + "; ".join(extra))
    r[0] += 1
    linea("PASO 6 – Matrices [A]", T)
    for rb in res.barras:
        matriz(f"[A]{rb.id}", rb.A, "0.0000", rb.filas, Dn)
    linea("PASO 7 – Matrices [k]", T)
    for rb in res.barras:
        matriz(f"[k]{rb.id}", rb.k, "0.00", rb.filas, rb.filas)
    linea("PASO 8 – [K]e = [A]ᵀ[k][A]", T)
    for rb in res.barras:
        matriz(f"[K]{rb.id}", rb.Kg, "0.00", Dn, Dn)
    linea("PASO 9 – [K]", T)
    matriz("[K]", res.K, "0.00", Dn, Dn)
    linea("PASO 10 – {r} locales", T)
    for rb in res.barras:
        matriz(f"{{r}}{rb.id}", rb.r, "0.0000", rb.filas)
    linea("PASO 11 – {R} = Σ[A]ᵀ{r} − {fuerzas aplicadas directamente en los nudos}", T)
    matriz("Σ[A]ᵀ{r}", res.Rs, "0.0000", Dn)
    matriz("Fuerzas en nudos", res.Q, "0.0000", Dn)
    matriz("{R}", res.R, "0.0000", Dn)
    linea("PASO 12 – {F} = −{R}", T)
    matriz("{F}", res.F, "0.0000", Dn)
    linea("PASO 13 – {D}", T)
    matriz("[K]⁻¹", np.linalg.inv(res.K), "0.0000000", Dn, Dn)
    matriz("{D}", res.D, "0.000000", Dn)
    linea("PASO 14 – Fuerzas internas", T)
    for rb in res.barras:
        ws.cell(r[0], 1, f"Barra {rb.id}").font = B
        r[0] += 1
        for k, (nom, v, fm) in enumerate((("{d}", rb.d, "0.000000"), ("{q}c", rb.qc, "0.0000"),
                                          ("{r}", rb.r, "0.0000"), ("{q} final", rb.q, "0.00"))):
            ws.cell(r[0], 2 + 2 * k, nom).font = B
            for i, val in enumerate(v):
                c = ws.cell(r[0] + 1 + i, 2 + 2 * k, float(val))
                c.number_format = fm
                c.fill = F
                c.border = box
        for i, fn in enumerate(rb.filas):
            ws.cell(r[0] + 1 + i, 1, fn).font = Font(italic=True, color="888888")
        r[0] += len(rb.filas) + 2
    ws.column_dimensions["A"].width = 16
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 13

    w2 = wb.create_sheet("Resultados")
    w2.append(["Reacciones", "Rx (t)", "Ry (t)", "M (t·m)"])
    for nid, v in res.reacciones.items():
        w2.append([f"Nudo {nid}", *[round(float(x), 4) for x in v]])
    w2.append([])
    w2.append(["Desplazamientos", "u (m)", "v (m)", "θ (rad)"])
    for nid, v in res.u_nudos.items():
        w2.append([f"Nudo {nid}", *[round(float(x), 7) for x in v]])
    w2.append([])
    w2.append(["Barra", "x (m)", "N (t)", "V (t)", "M (t·m)"])
    for rb in res.barras:
        xs, N, V, M, _, _ = esfuerzos(rb, 21)
        for x, n_, v_, m_ in zip(xs, N, V, M):
            w2.append([rb.id, round(float(x), 4), round(float(n_), 4), round(float(v_), 4), round(float(m_), 4)])
    for row in (1, 2):
        pass
    for c in w2[1]:
        c.font = B
    w2.column_dimensions["A"].width = 16
    if figuras:
        wf = wb.create_sheet("Figuras del método", 1)
        fila_img = 1
        lista = [("Pasos 2 a 5 – Numeración de nudos, barras y GDL (globales y locales)", figuras.get("numeracion"))]
        lista += [(f"Paso 6 – Configuración deformada D{j + 1} = 1", u)
                  for j, u in enumerate(figuras.get("configuraciones") or [])]
        lista += [("Paso 10 – Cargas de fijación {r} (ejes locales)", figuras.get("fijacion"))]
        for nombre, ruta_img in lista:
            if not ruta_img:
                continue
            wf.cell(fila_img, 1, nombre).font = T
            img = XImg(ruta_img)
            esc = 640 / img.width
            img.width, img.height = 640, int(img.height * esc)
            wf.add_image(img, f"A{fila_img + 1}")
            fila_img += int(img.height / 20) + 3
    if imagenes:
        w3 = wb.create_sheet("Diagramas")
        fila_img = 1
        for nombre, ruta_img in imagenes.items():
            w3.cell(fila_img, 1, nombre).font = T
            img = XImg(ruta_img)
            img.width, img.height = 640, 400
            w3.add_image(img, f"A{fila_img + 1}")
            fila_img += 23
    wb.save(ruta)
